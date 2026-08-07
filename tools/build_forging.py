#!/usr/bin/env python3
"""解析《橙装锻造》Sheet，生成 data/forging.js。

用法：
    python tools/build_forging.py [Excel路径]
"""
import json
import os
import re
import sys
from datetime import date

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(ROOT, "秦时相关（更新贯侯钟离昧）20260618.xlsx")
OUT_JS = os.path.join(ROOT, "data", "forging.js")

# 四大类区块：行号范围（含）与主装备名列
SECTIONS = [
    ("武器", 9, 68),
    ("盔甲", 72, 100),
    ("首饰", 104, 130),
    ("典籍", 134, 173),
]
STAGE_NAMES = [
    "0-1锻", "1-2锻", "2-3锻", "3-4锻", "4-5锻", "5-6锻",
    "6-7锻", "7-8锻", "8-9锻", "9-10锻", "10锻-红金",
]
STAGE_COLS = list(range(3, 14))  # C..M
SPLIT_RE = re.compile(r"[、,，]")


def cell_quality(cell):
    """按单元格字体颜色判定品质：紫色字体 → 紫；橙色/默认 → 橙。"""
    color = cell.font.color
    if color is None:
        return "橙"
    try:
        if color.type == "theme":
            # 本工作簿 theme:9 = accent6 F79646（橙色）
            return "橙"
        if color.type == "rgb" and color.rgb:
            s = str(color.rgb)
            if len(s) == 8:
                s = s[2:]
            r = int(s[0:2], 16)
            b = int(s[4:6], 16)
            return "紫" if b > r else "橙"
    except Exception:
        pass
    return "橙"


def parse_tokens(raw, quality):
    if raw is None:
        return [{"dash": True}]
    s = str(raw).strip()
    if not s or s == "-":
        return [{"dash": True}]
    return [{"n": part.strip(), "q": quality} for part in SPLIT_RE.split(s) if part.strip()]


def parse_sheet(path):
    """解析《橙装锻造》，返回 (summary, items)。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["橙装锻造"]
    summary = []
    for r in range(3, 7):
        cat = str(ws.cell(row=r, column=2).value).strip()
        summary.append({
            "cat": cat,
            "stages": [
                str(ws.cell(row=r, column=c).value).strip()
                for c in STAGE_COLS
            ],
            "total": str(ws.cell(row=r, column=14).value).strip(),
        })
    items = []
    for cat, start, end in SECTIONS:
        for r in range(start, end + 1):
            name_cell = ws.cell(row=r, column=2)
            raw_name = name_cell.value
            if raw_name is None:
                continue
            name = str(raw_name).strip()
            if not name:
                continue
            stages = []
            for idx, c in enumerate(STAGE_COLS):
                cell = ws.cell(row=r, column=c)
                stages.append({
                    "stage": STAGE_NAMES[idx],
                    "tokens": parse_tokens(cell.value, cell_quality(cell)),
                })
            items.append({
                "id": f"f-{len(items) + 1:04d}",
                "cat": cat,
                "name": name,
                "quality": cell_quality(name_cell),
                "stages": stages,
            })
    return summary, items


def build_data(path, out_path):
    summary, items = parse_sheet(path)
    m = re.search(r"(\d{8})", os.path.basename(path))
    payload = {
        "meta": {
            "sourceFile": os.path.basename(path),
            "sourceSheet": "橙装锻造",
            "version": m.group(1) if m else "unknown",
            "generatedAt": date.today().isoformat(),
            "total": len(items),
            "summaryTotal": len(summary),
            "stageNames": STAGE_NAMES,
        },
        "summary": summary,
        "items": items,
    }
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    js = (
        "/* 由 tools/build_forging.py 自动生成，请勿手改 */\n"
        "window.FORGING_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n"
    )
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(js)
    from collections import Counter
    counts = Counter(i["cat"] for i in items)
    print(f"生成完成：{out_path}")
    print(f"主装备总数：{len(items)}")
    for cat in ("武器", "盔甲", "首饰", "典籍"):
        print(f"  {cat}: {counts.get(cat, 0)}")
    return payload


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    build_data(xlsx, OUT_JS)
