#!/usr/bin/env python3
"""解析“新兵法”工作表，生成浏览器使用的兵法数据。"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOOK_NAME = "秦时相关（更新贯侯钟离昧）20260618.xlsx"


def default_xlsx(root=ROOT):
    """返回仓库主目录的默认 Excel；隔离工作树不使用自己的副本。"""
    main_root = root.parent.parent if root.parent.name == ".worktrees" else root
    return main_root / BOOK_NAME


DEFAULT_XLSX = default_xlsx()
OUT_JS = ROOT / "data" / "tactics.js"

ORDER = ["风", "林", "火", "山", "阴", "雷"]
STANDARD_BLOCKS = {
    "forest": {"name": "林", "range": (2, 20), "start_col": 2},
    "mountain": {"name": "山", "range": (24, 42), "start_col": 2},
    "wind": {"name": "风", "range": (45, 63), "start_col": 2},
    "fire": {"name": "火", "range": (66, 84), "start_col": 2},
}
SPECIAL_BLOCKS = {
    "yin": {"name": "阴", "range": (2, 20), "start_col": 19},
    "thunder": {"name": "雷", "range": (24, 42), "start_col": 19},
}

STANDARD_MANTRAS = {
    "forest": (("ling", "灵"), ("chan", "禅"), ("tong", "统")),
    "mountain": (("lie", "裂"), ("qia", "洽"), ("tong", "统")),
    "wind": (("qi", "齐"), ("biao", "镖"), ("tong", "统")),
    "fire": (("xin", "心"), ("jie", "解"), ("tong", "统")),
}
PARENTHESIS_ATTRIBUTE = re.compile(r"[（(]([^）)]+)[）)]")
ATTRIBUTE_ALIASES = {"攻击": "攻", "防御": "防"}


def number(value, coordinate):
    """将 Excel 数值标准化为 JSON 友好的 int 或 float。"""
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise ValueError(f"数值无法识别：{coordinate}={value!r}")
    return int(value) if float(value).is_integer() else float(value)


def displayed_percent(value, coordinate):
    if isinstance(value, str) and value.endswith("%"):
        return float(value[:-1])
    if isinstance(value, (int, float)):
        return float(value) * 100
    raise ValueError(f"百分比无法识别：{coordinate}={value!r}")


def extreme_value(mantra_rank):
    return (mantra_rank + 1) * 10


def attribute_name(value, coordinate):
    if value is None:
        raise ValueError(f"属性名称缺失：{coordinate}")
    name = str(value).strip()
    if not name:
        raise ValueError(f"属性名称缺失：{coordinate}")
    return name


def mantra_attribute(value, coordinate):
    text = attribute_name(value, coordinate)
    match = PARENTHESIS_ATTRIBUTE.search(text)
    if match:
        text = match.group(1).strip()
    return ATTRIBUTE_ALIASES.get(text, text)


def parse_proficiency(value, coordinate):
    if not isinstance(value, str):
        raise ValueError(f"熟练度无法识别：{coordinate}={value!r}")
    parts = value.strip().split("-", 1)
    if len(parts) != 2:
        raise ValueError(f"熟练度无法识别：{coordinate}={value!r}")
    try:
        return {"min": float(parts[0]), "max": float(parts[1])}
    except ValueError as error:
        raise ValueError(f"熟练度无法识别：{coordinate}={value!r}") from error


def compact_text(value):
    return re.sub(r"\s+", "", str(value or ""))


def expect_text(ws, row, column, expected, tactic_name, field):
    cell = ws.cell(row, column)
    if compact_text(cell.value) != compact_text(expected):
        raise ValueError(
            f"{tactic_name}{field}不匹配：{cell.coordinate}={cell.value!r}，应为{expected!r}"
        )


def expect_one_of(ws, row, column, expected_values, tactic_name, field):
    cell = ws.cell(row, column)
    actual = compact_text(cell.value)
    expected_compact = [compact_text(value) for value in expected_values]
    if actual not in expected_compact:
        raise ValueError(
            f"{tactic_name}{field}不匹配：{cell.coordinate}={cell.value!r}，"
            f"应为{'或'.join(repr(value) for value in expected_values)}"
        )


def expect_mantra_name(ws, row, column, expected, tactic_name):
    cell = ws.cell(row, column)
    text = str(cell.value or "").strip()
    actual = re.split(r"[（(]", text, maxsplit=1)[0].strip()
    if actual != expected:
        raise ValueError(
            f"{tactic_name}{expected}真言名称不匹配：{cell.coordinate}={cell.value!r}"
        )


def validate_standard_structure(ws, tactic_id, block):
    start_row, _ = block["range"]
    start_col = block["start_col"]
    tactic_name = f"{block['name']}兵法"
    group_row = start_row + 1
    header_row = start_row + 2
    expect_text(ws, start_row, start_col, tactic_name, tactic_name, "数据块标题")
    for offset, expected, field in (
        (0, "阶", "阶列表头"),
        (1, "基础属性", "基础属性表头"),
        (5, "熟练度", "熟练度表头"),
        (6, "进阶条件", "进阶材料表头"),
        (8, "演练单次需要号角", "单次演练号角表头"),
        (9, "每级需要号角", "保底号角表头"),
        (10, "真言（%）", "真言属性表头"),
        (14, "每级需要碎片", "真言碎片材料列"),
    ):
        expect_text(ws, group_row, start_col + offset, expected, tactic_name, field)
    mark_headers = [f"{block['name']}之印记"]
    if tactic_id == "fire":
        mark_headers.append("林之印记")
    expect_one_of(ws, header_row, start_col + 6, mark_headers, tactic_name, "进阶印记材料列")
    expect_text(ws, header_row, start_col + 7, "功勋", tactic_name, "功勋材料列")
    expected_mantras = [name for _, name in STANDARD_MANTRAS[tactic_id]] + ["极"]
    for offset, mantra_name in enumerate(expected_mantras, start=10):
        expect_mantra_name(ws, header_row, start_col + offset, mantra_name, tactic_name)


def validate_special_structure(ws, tactic_id, block):
    start_row, _ = block["range"]
    start_col = block["start_col"]
    tactic_name = f"{block['name']}兵法"
    group_row = start_row + 1
    header_row = start_row + 2
    expect_text(ws, start_row, start_col, tactic_name, tactic_name, "数据块标题")
    for offset, expected, field in (
        (0, "阶", "阶列表头"),
        (1, "基础属性", "基础属性表头"),
        (5, "额外属性", "额外属性表头"),
        (7, "进阶条件", "进阶材料表头"),
        (10, "真言（%）", "真言属性表头"),
    ):
        expect_text(ws, group_row, start_col + offset, expected, tactic_name, field)
    expect_text(ws, header_row, start_col + 7, f"{block['name']}之印记", tactic_name, "进阶印记材料列")
    expect_text(ws, header_row, start_col + 8, "功勋", tactic_name, "功勋材料列")
    expect_text(ws, header_row, start_col + 9, "号角", tactic_name, "号角材料列")
    expect_mantra_name(
        ws,
        header_row,
        start_col + 11,
        "殇" if tactic_id == "yin" else "盛",
        tactic_name,
    )


def required_number(ws, row, column, context):
    cell = ws.cell(row, column)
    if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
        raise ValueError(f"{context}缺失：{cell.coordinate}")
    try:
        return number(cell.value, cell.coordinate)
    except ValueError as error:
        raise ValueError(f"{context}数值无法识别：{cell.coordinate}={cell.value!r}") from error


def optional_number_or_zero(ws, row, column):
    cell = ws.cell(row, column)
    return 0 if cell.value is None else number(cell.value, cell.coordinate)


def required_percent(ws, row, column, context):
    cell = ws.cell(row, column)
    if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
        raise ValueError(f"{context}缺失：{cell.coordinate}")
    try:
        return displayed_percent(cell.value, cell.coordinate)
    except ValueError as error:
        raise ValueError(f"{context}百分比无法识别：{cell.coordinate}={cell.value!r}") from error


def required_proficiency(ws, row, column, context):
    cell = ws.cell(row, column)
    if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
        raise ValueError(f"{context}缺失：{cell.coordinate}")
    try:
        return parse_proficiency(cell.value, cell.coordinate)
    except ValueError as error:
        raise ValueError(f"{context}无法识别：{cell.coordinate}={cell.value!r}") from error


def parse_standard_block(ws, tactic_id, block):
    start_row, end_row = block["range"]
    start_col = block["start_col"]
    name = block["name"]
    tactic_name = f"{name}兵法"
    header_row = start_row + 2
    data_start = start_row + 3
    validate_standard_structure(ws, tactic_id, block)
    base_names = [
        attribute_name(ws.cell(header_row, start_col + offset).value, ws.cell(header_row, start_col + offset).coordinate)
        for offset in range(1, 5)
    ]

    ranks = []
    for row in range(data_start, end_row + 1):
        rank = required_number(ws, row, start_col, f"{tactic_name}阶数")
        base_attributes = []
        for offset, base_name in enumerate(base_names, start=1):
            value = ws.cell(row, start_col + offset).value
            if value is not None:
                base_attributes.append({
                    "name": base_name,
                    "value": number(value, ws.cell(row, start_col + offset).coordinate),
                    "unit": "flat",
                })
        ranks.append({
            "rank": rank,
            "baseAttributes": base_attributes,
            "extraAttributes": [],
            "proficiency": required_proficiency(ws, row, start_col + 5, f"{tactic_name}熟练度"),
            "advance": {
                "mark": optional_number_or_zero(ws, row, start_col + 6) if rank == 0 else required_number(
                    ws, row, start_col + 6, f"{tactic_name}{rank}阶进阶印记"
                ),
                "merit": optional_number_or_zero(ws, row, start_col + 7) if rank == 0 else required_number(
                    ws, row, start_col + 7, f"{tactic_name}{rank}阶功勋"
                ),
                "horn": 0,
            },
            "rehearsal": {
                "singleHorn": required_number(ws, row, start_col + 8, f"{tactic_name}{rank}阶单次演练号角"),
                "guaranteeHorn": required_number(ws, row, start_col + 9, f"{tactic_name}{rank}阶保底号角"),
            },
        })

    mantras = []
    for offset, (mantra_id, mantra_name) in enumerate(STANDARD_MANTRAS[tactic_id], start=10):
        header = ws.cell(header_row, start_col + offset)
        mantras.append({
            "id": mantra_id,
            "name": mantra_name,
            "attribute": mantra_attribute(header.value, header.coordinate),
            "unit": "flat" if mantra_name == "统" else "percent",
            "materialName": f"{mantra_name}真言碎片",
            "unlockTacticRank": 0,
            "maxRank": 9,
            "stages": [
                {
                    "rank": tactic_rank,
                    "tacticRank": tactic_rank,
                    "value": required_number(
                        ws,
                        data_start + tactic_rank,
                        start_col + offset,
                        f"{tactic_name}{mantra_name}真言{tactic_rank}阶属性",
                    ),
                    "fragments": required_number(
                        ws,
                        data_start + tactic_rank,
                        start_col + 14,
                        f"{tactic_name}{mantra_name}真言碎片",
                    ),
                }
                for tactic_rank in range(10)
            ],
        })

    mantras.append({
        "id": "extreme",
        "name": "极",
        "attribute": mantra_attribute(
            ws.cell(header_row, start_col + 13).value,
            ws.cell(header_row, start_col + 13).coordinate,
        ),
        "unit": "percent",
        "materialName": "极真言碎片",
        "unlockTacticRank": 10,
        "maxRank": 5,
        "stages": [
            {
                "rank": mantra_rank,
                "tacticRank": mantra_rank + 10,
                "value": extreme_value(mantra_rank),
                "fragments": required_number(
                    ws,
                    data_start + mantra_rank + 10,
                    start_col + 14,
                    f"{tactic_name}极真言碎片",
                ),
            }
            for mantra_rank in range(6)
        ],
    })
    return {
        "id": tactic_id,
        "name": f"{name}兵法",
        "kind": "standard",
        "markName": f"{name}之印记",
        "ranks": ranks,
        "mantras": mantras,
    }


def parse_special_block(ws, tactic_id, block):
    start_row, end_row = block["range"]
    start_col = block["start_col"]
    name = block["name"]
    tactic_name = f"{name}兵法"
    header_row = start_row + 2
    data_start = start_row + 3
    validate_special_structure(ws, tactic_id, block)
    base_names = [
        attribute_name(ws.cell(header_row, start_col + offset).value, ws.cell(header_row, start_col + offset).coordinate)
        for offset in range(1, 5)
    ]
    shield_name = attribute_name(ws.cell(header_row, start_col + 5).value, ws.cell(header_row, start_col + 5).coordinate)
    extra_percent_name = attribute_name(ws.cell(header_row, start_col + 6).value, ws.cell(header_row, start_col + 6).coordinate)

    ranks = []
    for row in range(data_start, end_row + 1):
        rank = required_number(ws, row, start_col, f"{tactic_name}阶数")
        base_value = required_percent(ws, row, start_col + 1, f"{tactic_name}{rank}阶组合基础属性")
        base_attributes = [
            {"name": base_name, "value": number(base_value, ws.cell(row, start_col + 1).coordinate), "unit": "percent"}
            for base_name in base_names
        ]
        ranks.append({
            "rank": rank,
            "baseAttributes": base_attributes,
            "extraAttributes": [
                {
                    "name": shield_name,
                    "value": required_number(ws, row, start_col + 5, f"{tactic_name}{rank}阶护盾"),
                    "unit": "flat",
                },
                {
                    "name": extra_percent_name,
                    "value": number(
                        required_percent(ws, row, start_col + 6, f"{tactic_name}{rank}阶额外属性"),
                        ws.cell(row, start_col + 6).coordinate,
                    ),
                    "unit": "percent",
                },
            ],
            "advance": {
                "mark": optional_number_or_zero(ws, row, start_col + 7) if rank == 0 else required_number(
                    ws, row, start_col + 7, f"{tactic_name}{rank}阶进阶印记"
                ),
                "merit": optional_number_or_zero(ws, row, start_col + 8) if rank == 0 else required_number(
                    ws, row, start_col + 8, f"{tactic_name}{rank}阶功勋"
                ),
                "horn": optional_number_or_zero(ws, row, start_col + 9) if rank == 0 else required_number(
                    ws, row, start_col + 9, f"{tactic_name}{rank}阶进阶号角"
                ),
            },
        })

    mantra_name = attribute_name(ws.cell(header_row, start_col + 11).value, ws.cell(header_row, start_col + 11).coordinate)
    mantra_attribute_name = attribute_name(ws.cell(header_row, start_col + 10).value, ws.cell(header_row, start_col + 10).coordinate)
    return {
        "id": tactic_id,
        "name": f"{name}兵法",
        "kind": "special",
        "markName": f"{name}之印记",
        "ranks": ranks,
        "mantras": [{
            "id": "shang" if tactic_id == "yin" else "sheng",
            "name": mantra_name,
            "attribute": mantra_attribute_name,
            "unit": "percent",
            "materialName": f"{mantra_name}真言碎片",
            "unlockTacticRank": 0,
            "maxRank": 15,
            "stages": [
                {
                    "rank": tactic_rank,
                    "tacticRank": tactic_rank,
                    "value": number(
                        required_percent(
                            ws,
                            data_start + tactic_rank,
                            start_col + 10,
                            f"{tactic_name}{mantra_name}真言{tactic_rank}阶属性",
                        ),
                        ws.cell(data_start + tactic_rank, start_col + 10).coordinate,
                    ),
                    "fragments": required_number(
                        ws,
                        data_start + tactic_rank,
                        start_col + 11,
                        f"{tactic_name}{mantra_name}真言碎片",
                    ),
                }
                for tactic_rank in range(16)
            ],
        }],
    }


def validate_payload(payload):
    items = payload["items"]
    if [item["name"][0] for item in items] != ORDER:
        raise ValueError("兵法顺序或数据块缺失")
    for item in items:
        ranks = [row["rank"] for row in item["ranks"]]
        if ranks != list(range(16)):
            raise ValueError(f"{item['name']}阶数不连续：{ranks}")
        for mantra in item["mantras"]:
            expected = list(range(mantra["maxRank"] + 1))
            actual = [stage["rank"] for stage in mantra["stages"]]
            if actual != expected:
                raise ValueError(f"{item['name']}·{mantra['name']}真言阶数不连续：{actual}")


def parse_worksheet(sheet):
    parsed = {
        "wind": parse_standard_block(sheet, "wind", STANDARD_BLOCKS["wind"]),
        "forest": parse_standard_block(sheet, "forest", STANDARD_BLOCKS["forest"]),
        "fire": parse_standard_block(sheet, "fire", STANDARD_BLOCKS["fire"]),
        "mountain": parse_standard_block(sheet, "mountain", STANDARD_BLOCKS["mountain"]),
        "yin": parse_special_block(sheet, "yin", SPECIAL_BLOCKS["yin"]),
        "thunder": parse_special_block(sheet, "thunder", SPECIAL_BLOCKS["thunder"]),
    }
    payload = {"meta": {"order": ORDER}, "items": [parsed[key] for key in ("wind", "forest", "fire", "mountain", "yin", "thunder")]}
    validate_payload(payload)
    return payload


def parse_sheet(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        return parse_worksheet(workbook["新兵法"])
    finally:
        workbook.close()


def build_data(path=DEFAULT_XLSX, out_path=OUT_JS):
    payload = parse_sheet(path)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(
        "/* 由 tools/build_tactics.py 自动生成，请勿手改 */\n"
        "window.TACTICS_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    print(f"generated {len(payload['items'])} tactics from {path}")
    return payload


if __name__ == "__main__":
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    build_data(source)
