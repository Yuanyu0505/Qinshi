#!/usr/bin/env python3
"""读取“战匣丹囊”工作表，生成浏览器与节点测试共用的静态数据。"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOOK_NAME = "秦时相关（更新贯侯钟离昧）20260618.xlsx"
SHEET_NAME = "战匣丹囊"
OUT_JS = ROOT / "data" / "battle-box-pill-pouch.js"


def default_xlsx(root=ROOT):
    """隔离工作树默认读取主工作区中的最新工作簿。"""
    main_root = root.parent.parent if root.parent.name == ".worktrees" else root
    return main_root / BOOK_NAME


def integer(value, coordinate):
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise ValueError(f"数值无法识别：{coordinate}={value!r}")
    number = int(value)
    if float(value) != number:
        raise ValueError(f"预期整数：{coordinate}={value!r}")
    return number


def build(source):
    workbook = load_workbook(source, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"找不到工作表：{SHEET_NAME}")
    sheet = workbook[SHEET_NAME]

    battle_levels = [{
        "level": 0,
        "pearls": 0,
        "shells": 0,
        "attack": 0,
        "defense": 0,
        "health": 0,
        "pvpMitigation": 0,
    }]
    for row in range(21, 111):
        level = integer(sheet.cell(row, 2).value, f"B{row}")
        battle_levels.append({
            "level": level,
            "pearls": integer(sheet.cell(row, 7).value, f"G{row}"),
            "shells": integer(sheet.cell(row, 8).value, f"H{row}"),
            "attack": integer(sheet.cell(row, 3).value, f"C{row}"),
            "defense": integer(sheet.cell(row, 4).value, f"D{row}"),
            "health": 557660 if level == 61 else integer(sheet.cell(row, 5).value, f"E{row}"),
            "pvpMitigation": integer(sheet.cell(row, 6).value, f"F{row}"),
        })

    pouch_levels = [{
        "level": 0,
        "pearls": 0,
        "shells": 0,
        "bonusPercent": 0,
    }]
    for row in range(116, 206):
        level = integer(sheet.cell(row, 2).value, f"B{row}")
        pouch_levels.append({
            "level": level,
            "pearls": integer(sheet.cell(row, 3).value, f"C{row}"),
            "shells": integer(sheet.cell(row, 4).value, f"D{row}"),
            "bonusPercent": integer(sheet.cell(row, 5).value, f"E{row}"),
        })

    if [item["level"] for item in battle_levels] != list(range(91)):
        raise ValueError("战匣等级必须连续覆盖 0–90 级。")
    if [item["level"] for item in pouch_levels] != list(range(91)):
        raise ValueError("丹囊等级必须连续覆盖 0–90 级。")

    return {
        "meta": {
            "sourceFile": BOOK_NAME,
            "sourceSheet": SHEET_NAME,
            "version": "20260618",
            "generatedAt": "2026-09-02",
            "materialMeaning": "上一等级升至本等级所需",
            "slotUnlockNotice": "槽位分级解锁规则待准确数据补充",
        },
        "equipmentSlots": [
            {"id": "weapon", "name": "武器", "categories": ["武器", "神兵武器"]},
            {"id": "armor", "name": "防具", "categories": ["防具", "神兵防具"]},
            {"id": "accessory", "name": "饰品", "categories": ["饰品", "神兵饰品"]},
            {"id": "book", "name": "典籍", "categories": ["典籍", "神兵典籍"]},
        ],
        "battleQualityCaps": {
            "green": 2,
            "blue": 3,
            "purple": 5,
            "orange": 8,
            "orangeGold": 12,
            "red": 15,
            "redGold": 20,
        },
        "battleQualityNames": {
            "green": "绿色",
            "blue": "蓝色",
            "purple": "紫色",
            "orange": "橙色",
            "orangeGold": "橙金",
            "red": "红色",
            "redGold": "红金",
        },
        "pouchQualityCaps": {
            "green": 1,
            "blue": 2,
            "purple": 3,
            "orange": 4,
            "heaven": 5,
            "immortal": 6,
            "sacred": 8,
            "divine": 10,
        },
        "pouchQualityNames": {
            "green": "绿色",
            "blue": "蓝色",
            "purple": "紫色",
            "orange": "橙色",
            "heaven": "天级",
            "immortal": "仙级",
            "sacred": "圣级",
            "divine": "神级",
        },
        "battlePlayerCaps": [
            {"min": 46, "max": 53, "cap": 70},
            {"min": 54, "max": None, "cap": 90},
        ],
        "pouchPlayerCaps": [
            {"min": 46, "max": 56, "cap": 40},
            {"min": 57, "max": 59, "cap": 50},
            {"min": 60, "max": 66, "cap": 60},
            {"min": 67, "max": 73, "cap": 70},
            {"min": 74, "max": 79, "cap": 80},
            {"min": 80, "max": None, "cap": 90},
        ],
        "defaults": {
            "baseCap": 10,
            "defaultQuality": "orange",
            "purchase": {
                "pearls": {"packSize": 5, "packPrice": 20},
                "shells": {"packSize": 10, "packPrice": 300},
            },
        },
        "battleLevels": battle_levels,
        "pouchLevels": pouch_levels,
    }


def write_js(data, output=OUT_JS):
    payload = json.dumps(data, ensure_ascii=False, indent=2)
    content = (
        "/* 由 tools/build_battle_box_pill_pouch.py 自动生成，请勿手改 */\n"
        "(function (root, factory) {\n"
        "  var data = factory();\n"
        "  if (typeof module === \"object\" && module.exports) module.exports = data;\n"
        "  if (root) root.BATTLE_BOX_PILL_POUCH_DATA = data;\n"
        "})(typeof globalThis !== \"undefined\" ? globalThis : this, function () {\n"
        "  \"use strict\";\n"
        f"  return {payload};\n"
        "});\n"
    )
    output.write_text(content, encoding="utf-8")


def main(argv):
    source = Path(argv[1]).resolve() if len(argv) > 1 else default_xlsx()
    if not source.exists():
        raise FileNotFoundError(source)
    write_js(build(source))
    print(f"generated {OUT_JS.relative_to(ROOT)} from {source}")


if __name__ == "__main__":
    main(sys.argv)
