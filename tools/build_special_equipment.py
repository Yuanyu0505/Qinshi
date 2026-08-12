#!/usr/bin/env python3
"""解析《特殊属性装备》《典籍属性》Sheet，生成 data/special-equipment.js。

用法：
    python tools/build_special_equipment.py [Excel路径]
默认使用项目根目录下的《秦时相关（更新贯侯钟离昧）20260618.xlsx》。
"""
import json
import os
import re
import sys
from collections import Counter
from datetime import date

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(ROOT, "秦时相关（更新贯侯钟离昧）20260618.xlsx")
OUT_JS = os.path.join(ROOT, "data", "special-equipment.js")

# 每个区块 10 列：名称、主属性、4 档 × 2 列（1-based 列号）
BLOCKS = {
    "武器": {"name": 2, "main": 3, "tiers": {"橙色": [4, 5], "橙金": [6, 7], "红色": [8, 9], "红金": [10, 11]}},
    "防具": {"name": 13, "main": 14, "tiers": {"橙色": [15, 16], "橙金": [17, 18], "红色": [19, 20], "红金": [21, 22]}},
    "饰品": {"name": 24, "main": 25, "tiers": {"橙色": [26, 27], "橙金": [28, 29], "红色": [30, 31], "红金": [32, 33]}},
}
HEADER_NAMES = {"武器名称", "防具名称", "饰品名称", "神兵武器名称", "神兵防具名称", "神兵饰品名称"}
ATTR_PAT = re.compile(r"^(\d+(?:\.\d+)?)%\s*(.+)$")
NORMALIZE = {"血量": "血", "减伤": "技伤减免", "减免": "技伤减免"}
STATUSES = {"无", "暂未开放"}
ATTR_ORDER = [
    "攻", "血", "防", "攻防血", "穿透", "暴击", "暴伤", "技伤减免", "抗暴",
    "速", "闪避", "招架", "敌方减攻", "敌方减防", "敌方减血",
]
CATEGORY_ORDER = ["武器", "防具", "饰品", "典籍", "神兵武器", "神兵防具", "神兵饰品", "神兵典籍"]
BOOK_TIERS = {
    "紫色": [(0, 3), (5, 4), (10, 5)],
    "橙色": [(0, 6), (5, 7), (10, 8)],
    "橙金": [(0, 9), (5, 10), (10, 11), (15, 12)],
    "红色": [(0, 13), (5, 14), (10, 15), (15, 16)],
    "红金": [(0, 17), (5, 18), (10, 19), (15, 20)],
}
SPEED_PAT = re.compile(r"^(\d+(?:\.\d+)?)\s*速$")
ENEMY_PAT = re.compile(r"^敌方-\s*(\d+(?:\.\d+)?)%\s*(攻|防|血)$")


def parse_cell(value):
    """单个单元格 → token dict / None。"""
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if s in STATUSES:
        return {"s": s}
    m = ATTR_PAT.match(s)
    if not m:
        return {"raw": s}
    original_type = m.group(2).strip()
    t = NORMALIZE.get(original_type, original_type)
    raw = f"{m.group(1)}%{t}" if original_type in {"减伤", "减免"} else s
    return {"t": t, "v": float(m.group(1)), "raw": raw}


def parse_book_cell(value):
    """典籍单元格 → token 列表；支持多行、速度、敌方减属性和复合攻防血。"""
    if value is None:
        return []
    text = str(value).strip()
    if not text or text == "-":
        return []
    tokens = []
    for line in re.split(r"\r?\n", text):
        s = line.strip()
        if not s or s == "-":
            continue
        speed = SPEED_PAT.match(s)
        if speed:
            tokens.append({"t": "速", "v": float(speed.group(1)), "raw": f"{speed.group(1)}速"})
            continue
        enemy = ENEMY_PAT.match(s)
        if enemy:
            stat = enemy.group(2)
            tokens.append({"t": f"敌方减{stat}", "v": float(enemy.group(1)), "raw": f"敌方-{enemy.group(1)}%{stat}"})
            continue
        token = parse_cell(s)
        if token is None:
            continue
        if "raw" in token and "t" not in token:
            tokens.append(token)
            continue
        if token.get("t") == "攻防血":
            token["matches"] = ["攻", "防", "血", "攻防血"]
        tokens.append(token)
    return tokens


def update_max_map(max_map, token):
    """将 token 数值写入自身类型及复合匹配类型的最高值表。"""
    if "t" not in token or not isinstance(token.get("v"), (int, float)):
        return
    for attr in token.get("matches", [token["t"]]):
        max_map[attr] = max(max_map.get(attr, 0), token["v"])


def parse_book_sheet(wb):
    """解析“典籍属性”，返回 (items, anomalies)。"""
    ws = wb["典籍属性"]
    items = []
    anomalies = []
    for row_number in range(2, 54):
        name = str(ws.cell(row_number, 1).value or "").strip()
        if not name:
            continue
        main = str(ws.cell(row_number, 2).value or "").strip()
        main_key = main.split("、", 1)[0].strip()
        if row_number <= 19:
            book_group = "初始紫色典籍"
            category = "典籍"
            tiers_to_read = BOOK_TIERS
        elif row_number <= 40:
            book_group = "初始橙色典籍"
            category = "典籍"
            tiers_to_read = {tier: columns for tier, columns in BOOK_TIERS.items() if tier != "紫色"}
        else:
            book_group = "神兵典籍"
            category = "神兵典籍"
            tiers_to_read = {tier: columns for tier, columns in BOOK_TIERS.items() if tier != "紫色"}

        stages = {}
        tiers = {}
        max_map = {}
        for tier, stage_columns in tiers_to_read.items():
            stage_rows = []
            flat_tokens = []
            for stage, column in stage_columns:
                tokens = parse_book_cell(ws.cell(row_number, column).value)
                for token in tokens:
                    if "raw" in token and "t" not in token:
                        anomalies.append({
                            "sheet": "典籍属性", "cell": ws.cell(row_number, column).coordinate,
                            "cat": category, "name": name, "tier": tier, "value": token["raw"],
                        })
                    update_max_map(max_map, token)
                stage_rows.append({"stage": stage, "tokens": tokens})
                flat_tokens.extend(tokens)
            stages[tier] = stage_rows
            tiers[tier] = flat_tokens

        prefix = "db" if category == "神兵典籍" else "b"
        items.append({
            "id": f"{prefix}-{len(items) + 1:04d}",
            "cat": category,
            "name": name,
            "main": main,
            "mainKey": main_key,
            "bookGroup": book_group,
            "sourceOrder": row_number - 2,
            "tiers": tiers,
            "stages": stages,
            "max": max_map,
        })
    return items, anomalies


def parse_sheet(path):
    """解析工作簿，返回 (items, anomalies)。"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["特殊属性装备"]
    rows = list(ws.iter_rows(values_only=True))
    items = []
    anomalies = []
    for r in range(2, ws.max_row + 1):
        if r == 20:  # 神兵区块表头行
            continue
        row = rows[r - 1]
        for base, b in BLOCKS.items():
            raw_name = row[b["name"] - 1]
            if raw_name is None:
                continue
            name = str(raw_name).strip()
            if not name or name in HEADER_NAMES:
                continue
            main = str(row[b["main"] - 1]).strip() if row[b["main"] - 1] is not None else ""
            cat = "神兵" + base if r >= 21 else base
            tiers = {}
            max_map = {}
            for tier, cols in b["tiers"].items():
                tokens = []
                for c in cols:
                    tok = parse_cell(row[c - 1])
                    if tok is None:
                        continue
                    tokens.append(tok)
                    if "t" in tok:
                        update_max_map(max_map, tok)
                    elif "raw" in tok:
                        anomalies.append(
                            {"row": r, "cat": cat, "name": name, "tier": tier, "cell": tok["raw"]}
                        )
                tiers[tier] = tokens
            prefix = {"武器": "w", "防具": "f", "饰品": "s"}[base]
            items.append({
                "id": f"{prefix}-{len(items) + 1:04d}",
                "cat": cat,
                "name": name,
                "main": main,
                "tiers": tiers,
                "max": max_map,
            })
    book_items, book_anomalies = parse_book_sheet(wb)
    items.extend(book_items)
    anomalies.extend(book_anomalies)
    wb.close()
    return items, anomalies


def build_data(path, out_path):
    items, anomalies = parse_sheet(path)
    if anomalies:
        for a in anomalies:
            print("异常单元格:", a)
        raise SystemExit("存在无法解析的单元格，已中止生成。")
    observed = Counter()
    for item in items:
        observed.update(item["max"].keys())
    attr_types = [a for a in ATTR_ORDER if a in observed]
    attr_types += sorted(observed.keys() - set(ATTR_ORDER))
    m = re.search(r"(\d{8})", os.path.basename(path))
    payload = {
        "meta": {
            "sourceFile": os.path.basename(path),
            "sourceSheet": "特殊属性装备、典籍属性",
            "version": m.group(1) if m else "unknown",
            "generatedAt": date.today().isoformat(),
            "total": len(items),
            "attrTypes": attr_types,
            "categories": CATEGORY_ORDER,
        },
        "items": items,
    }
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    js = (
        "/* 由 tools/build_special_equipment.py 自动生成，请勿手改 */\n"
        "window.SPECIAL_EQUIPMENT_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n"
    )
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(js)
    counts = Counter(i["cat"] for i in items)
    print(f"生成完成：{out_path}")
    print(f"装备总数：{len(items)}")
    for cat in CATEGORY_ORDER:
        print(f"  {cat}: {counts.get(cat, 0)}")
    print("属性类型：", "、".join(attr_types))
    return payload


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    build_data(xlsx, OUT_JS)
