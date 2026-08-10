#!/usr/bin/env python3
"""解析《图鉴汇总》Sheet，生成 data/atlas.js。

用法：
    python tools/build_atlas.py [Excel路径]
"""
import json
import os
import re
import sys
from datetime import date

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(ROOT, "秦时相关（更新贯侯钟离昧）20260618.xlsx")
OUT_JS = os.path.join(ROOT, "data", "atlas.js")

# (图鉴名, 名称列, 数据起始行, 数据结束行)
BLOCKS = [
    ("攻", 4, 3, 52),
    ("血", 12, 3, 65),
    ("内力", 20, 3, 31),
    ("防", 28, 3, 37),
]
STAGES = [("5→6", 6), ("7→8", 8), ("9→10", 10)]
UPGRADE_START_ROW = 51
UPGRADE_END_ROW = 69
UPGRADE_COLS = (20, 21, 22, 23, 24)  # T..X


def cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def int_value(value):
    if value is None or value == "":
        return 0
    return int(value)


def cell_quality(cell):
    """按单元格底色判定品质：紫色填充(theme:7/8064A2) → 紫；橙色(theme:9/F79646) → 橙。"""
    fill = cell.fill
    if fill is None or fill.patternType is None:
        return "橙"
    color = fill.fgColor
    if color is None:
        return "橙"
    try:
        if color.type == "theme":
            return "紫" if color.theme == 7 else "橙"
        if color.type == "rgb" and color.rgb:
            s = str(color.rgb)
            if len(s) == 8:
                s = s[2:]
            r = int(s[0:2], 16)
            b = int(s[4:6], 16)
            return "紫" if b > r else "橙"
    except Exception:
        pass
    return "橙"


def parse_sheet(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["图鉴汇总"]
    rows = list(ws.iter_rows(values_only=True))
    items = []
    for atlas, name_col, start, end in BLOCKS:
        for r in range(start, end + 1):
            name = cell_text(rows[r - 1][name_col - 1])
            if not name:
                continue
            stages = []
            for idx, (key, stage_end) in enumerate(STAGES):
                raw = cell_text(rows[r - 1][name_col + idx])
                stages.append({
                    "key": key,
                    "end": stage_end,
                    "items": [
                        {"n": p.strip(), "q": cell_quality(ws.cell(row=r, column=name_col + 1 + idx))}
                        for p in raw.replace("、", ",").split(",") if p.strip()
                    ],
                })
            acquire = cell_text(rows[r - 1][name_col + 3])
            group = cell_text(rows[r - 1][name_col + 4])
            level_raw = cell_text(rows[r - 1][name_col + 5])
            level = int(level_raw) if level_raw.isdigit() else 0
            items.append({
                "id": f"t-{len(items) + 1:04d}",
                "atlas": atlas,
                "name": name,
                "stages": stages,
                "acquire": acquire,
                "group": group,
                "level": level,
            })
    return items


def parse_upgrade_stages(ws):
    """解析 T51:X69 图鉴升级成本表。"""
    stages = []
    for r in range(UPGRADE_START_ROW, UPGRADE_END_ROW + 1):
        raw = cell_text(ws.cell(row=r, column=UPGRADE_COLS[0]).value)
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", raw)
        if not m:
            raise ValueError(f"图鉴升级阶段格式无法识别：T{r}={raw!r}")
        start = int(m.group(1))
        end = int(m.group(2))
        if end != start + 1:
            raise ValueError(f"图鉴升级阶段必须连续：T{r}={raw!r}")
        knots = int_value(ws.cell(row=r, column=UPGRADE_COLS[1]).value)
        souls = int_value(ws.cell(row=r, column=UPGRADE_COLS[2]).value)
        equipment = int_value(ws.cell(row=r, column=UPGRADE_COLS[3]).value)
        growth = int_value(ws.cell(row=r, column=UPGRADE_COLS[4]).value)
        stages.append({
            "key": f"{start}→{end}",
            "from": start,
            "to": end,
            "knots": knots,
            "souls": souls,
            "needsEquipment": equipment > 0,
            "growth": growth,
        })
    return stages


def build_data(path, out_path):
    items = parse_sheet(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    upgrade_stages = parse_upgrade_stages(wb["图鉴汇总"])
    max_level = max(stage["to"] for stage in upgrade_stages)
    m = re.search(r"(\d{8})", os.path.basename(path))
    from collections import Counter
    counts = Counter(i["atlas"] for i in items)
    payload = {
        "meta": {
            "sourceFile": os.path.basename(path),
            "sourceSheet": "图鉴汇总",
            "version": m.group(1) if m else "unknown",
            "generatedAt": date.today().isoformat(),
            "total": len(items),
            "atlasOrder": ["攻", "血", "内力", "防"],
            "counts": dict(counts),
            "upgradeStages": upgrade_stages,
            "maxLevel": max_level,
            "defaultTargetLevel": min(max(item["level"] for item in items), max_level),
        },
        "items": items,
    }
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    js = (
        "/* 由 tools/build_atlas.py 自动生成，请勿手改 */\n"
        "window.ATLAS_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n"
    )
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(js)
    print(f"生成完成：{out_path}")
    print(f"图鉴弟子总数：{len(items)}")
    for atlas in ("攻", "血", "内力", "防"):
        print(f"  {atlas}图鉴：{counts.get(atlas, 0)}")
    return payload


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    build_data(xlsx, OUT_JS)
