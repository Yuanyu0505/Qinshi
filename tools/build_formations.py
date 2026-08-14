#!/usr/bin/env python3
"""解析“合阵”工作表，生成浏览器使用的合阵数据。"""

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOOK_NAME = "秦时相关（更新贯侯钟离昧）20260618.xlsx"
OUT_JS = ROOT / "data" / "formations.js"


def default_xlsx(root=ROOT):
    """隔离工作树始终读取主仓库中的最新工作簿。"""
    main_root = root.parent.parent if root.parent.name == ".worktrees" else root
    return main_root / BOOK_NAME


DEFAULT_XLSX = default_xlsx()

# 阵眼石、合阵名、弟子名、攻、血、防、转换规则、推荐位置、起始行、结束行
BLOCKS = (
    (2, 3, 4, 5, 6, 7, 8, 9, ((2, 8), (9, 14), (15, 20), (21, 26), (27, 32), (33, 39), (40, 47), (48, 59))),
    (11, 12, 13, 14, 15, 16, 17, 18, ((2, 9), (10, 17), (18, 25), (26, 33), (34, 40), (41, 48), (49, 56), (57, 63), (64, 71))),
    (21, 22, 23, 24, 25, 26, 27, 28, ((2, 9), (10, 17))),
)

SUMMARY_ROWS = range(77, 96)
RULE_RE = re.compile(
    r"(?P<rate>\d+(?:\.\d+)?)%\s*(?P<source>攻|血|防)(?:属性)?\s*[→\-]\s*"
    r"(?P<target>全体攻|全体血|全体防|全体内力|全体护盾|追加伤害)"
)


def compact_text(value):
    return re.sub(r"\s+", "", str(value or ""))


def normalize_name(value):
    return re.sub(r"[·•・\s]+", "", str(value or "").strip())


def required_text(cell, context):
    text = str(cell.value or "").strip()
    if not text:
        raise ValueError(f"{context}缺失：{cell.coordinate}")
    return text


def required_nonnegative_number(cell, context):
    value = cell.value
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0:
        raise ValueError(f"{context}必须为非负数字：{cell.coordinate}={value!r}")
    return int(value) if float(value).is_integer() else float(value)


def required_position(cell, context):
    value = required_nonnegative_number(cell, context)
    if not isinstance(value, int) or value < 1:
        raise ValueError(f"{context}必须为正整数：{cell.coordinate}={cell.value!r}")
    return value


def parse_rule(value, coordinate):
    raw = compact_text(value).replace("－", "-").replace("—", "-")
    match = RULE_RE.fullmatch(raw)
    if not match:
        raise ValueError(f"转换规则无法识别：{coordinate}={value!r}")
    rate = float(match.group("rate"))
    return {
        "sourceAttribute": match.group("source"),
        "ratePercent": int(rate) if rate.is_integer() else rate,
        "targetAttribute": match.group("target"),
        "rawRule": str(value).strip(),
        "sourceCell": coordinate,
    }


def parse_summary(sheet):
    summary = {}
    for row in SUMMARY_ROWS:
        stone_cell = sheet.cell(row, 10)
        name_cell = sheet.cell(row, 11)
        main_cell = sheet.cell(row, 12)
        stone = required_nonnegative_number(stone_cell, "阵眼石")
        name = required_text(name_cell, "合阵名称")
        normalized = normalize_name(name)
        if normalized in summary:
            raise ValueError(f"官方汇总存在重复合阵：{name_cell.coordinate}={name!r}")
        supports = []
        for position, column in enumerate(range(13, 19), start=1):
            value = sheet.cell(row, column).value
            if value is None or not str(value).strip() or normalize_name(value) == "无":
                continue
            supports.append({
                "position": position,
                "name": normalize_name(value),
                "rawName": str(value).strip(),
                "sourceCell": sheet.cell(row, column).coordinate,
            })
        summary[normalized] = {
            "stone": stone,
            "name": name,
            "officialMain": normalize_name(required_text(main_cell, f"{name}官方主将")),
            "officialMainRaw": str(main_cell.value).strip(),
            "supports": supports,
            "sourceRow": row,
        }
    return summary


def parse_block(sheet, columns, start_row, end_row):
    stone_col, name_col, disciple_col, attack_col, health_col, defense_col, rule_col, position_col = columns
    formation_name_cell = sheet.cell(start_row, name_col)
    raw_formation_name = required_text(formation_name_cell, "合阵名称")
    formation_name = normalize_name(raw_formation_name)
    stone = required_nonnegative_number(sheet.cell(start_row, stone_col), f"{formation_name}阵眼石")
    candidates = []
    slots_by_position = {}
    for row in range(start_row, end_row + 1):
        name_cell = sheet.cell(row, disciple_col)
        raw_name = required_text(name_cell, f"{formation_name}候选弟子")
        name = normalize_name(raw_name)
        candidate = {
            "id": f"disciple-{name}",
            "name": name,
            "rawName": raw_name,
            "level1": {
                "attack": required_nonnegative_number(sheet.cell(row, attack_col), f"{formation_name}·{name}攻"),
                "health": required_nonnegative_number(sheet.cell(row, health_col), f"{formation_name}·{name}血"),
                "defense": required_nonnegative_number(sheet.cell(row, defense_col), f"{formation_name}·{name}防"),
            },
            "officialPosition": None,
            "sourceRow": row,
        }
        position_cell = sheet.cell(row, position_col)
        if position_cell.value is not None and str(position_cell.value).strip():
            position = required_position(position_cell, f"{formation_name}·{name}推荐位置")
            if position in slots_by_position:
                raise ValueError(f"{formation_name}助阵{position}重复：{position_cell.coordinate}")
            rule_cell = sheet.cell(row, rule_col)
            slot = parse_rule(rule_cell.value, rule_cell.coordinate)
            slot.update({
                "position": position,
                "officialDisciple": name,
                "officialDiscipleRaw": raw_name,
            })
            slots_by_position[position] = slot
            candidate["officialPosition"] = position
        candidates.append(candidate)
    return {
        "id": f"formation-{formation_name}",
        "name": formation_name,
        "rawName": raw_formation_name,
        "stone": stone,
        "officialMain": None,
        "officialMainRaw": None,
        "slots": [slots_by_position[key] for key in sorted(slots_by_position)],
        "candidates": candidates,
        "sourceRange": f"{formation_name_cell.coordinate}:{sheet.cell(end_row, position_col).coordinate}",
    }


def validate_payload(payload):
    items = payload.get("items", [])
    if len(items) != 19:
        raise ValueError(f"合阵数量必须为19，实际为{len(items)}")
    candidate_count = sum(len(item.get("candidates", [])) for item in items)
    if candidate_count != 144:
        raise ValueError(f"候选弟子记录必须为144，实际为{candidate_count}")

    formation_names = set()
    global_candidates = set()
    for item in items:
        name = item["name"]
        if name in formation_names:
            raise ValueError(f"合阵名称重复：{name}")
        formation_names.add(name)
        slots = item["slots"]
        positions = [slot["position"] for slot in slots]
        if len(slots) not in (5, 6) or positions != list(range(1, len(slots) + 1)):
            raise ValueError(f"{name}助阵位置必须为连续的1至5或1至6，实际为{positions}")
        names = [candidate["name"] for candidate in item["candidates"]]
        if len(names) != len(set(names)):
            raise ValueError(f"{name}名称标准化后存在重复弟子")
        for candidate_name in names:
            if candidate_name in global_candidates:
                raise ValueError(f"弟子跨合阵重复：{candidate_name}")
            global_candidates.add(candidate_name)
        if item["officialMain"] not in set(names):
            raise ValueError(f"{name}官方主将不在候选名单：{item['officialMain']}")
        for slot in slots:
            if slot["officialDisciple"] not in set(names):
                raise ValueError(
                    f"{name}助阵{slot['position']}官方弟子不在候选名单：{slot['officialDisciple']}"
                )


def parse_worksheet(sheet):
    summary = parse_summary(sheet)
    items = []
    for block in BLOCKS:
        columns = block[:8]
        for start_row, end_row in block[8]:
            item = parse_block(sheet, columns, start_row, end_row)
            official = summary.get(item["name"])
            if not official:
                raise ValueError(f"{item['name']}未在J76:R95官方汇总中找到")
            if official["stone"] != item["stone"]:
                raise ValueError(
                    f"{item['name']}阵眼石不一致：资料区={item['stone']}，汇总区={official['stone']}"
                )
            item["officialMain"] = official["officialMain"]
            item["officialMainRaw"] = official["officialMainRaw"]
            summary_supports = {entry["position"]: entry for entry in official["supports"]}
            if len(summary_supports) != len(item["slots"]):
                raise ValueError(
                    f"{item['name']}官方助阵数量不一致：资料区={len(item['slots'])}，汇总区={len(summary_supports)}"
                )
            for slot in item["slots"]:
                expected = summary_supports.get(slot["position"])
                if not expected or expected["name"] != slot["officialDisciple"]:
                    coordinate = expected["sourceCell"] if expected else f"J{official['sourceRow']}:R{official['sourceRow']}"
                    raise ValueError(
                        f"{item['name']}助阵{slot['position']}官方弟子不一致："
                        f"资料区={slot['officialDisciple']}，汇总区={expected['name'] if expected else None}（{coordinate}）"
                    )
            items.append(item)
    payload = {
        "meta": {
            "sheet": "合阵",
            "formationCount": len(items),
            "candidateCount": sum(len(item["candidates"]) for item in items),
            "order": [item["id"] for item in items],
        },
        "items": items,
    }
    validate_payload(payload)
    return payload


def parse_sheet(path):
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if "合阵" not in workbook.sheetnames:
            raise ValueError("工作簿缺少“合阵”工作表")
        return parse_worksheet(workbook["合阵"])
    finally:
        workbook.close()


def write_output(payload, out_path=OUT_JS):
    validate_payload(payload)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    content = (
        "/* 由 tools/build_formations.py 自动生成，请勿手改 */\n"
        "window.FORMATIONS_DATA = "
        + json.dumps(payload, ensure_ascii=False, indent=2)
        + ";\n"
    )
    temporary = out_path.with_name(out_path.name + ".tmp")
    try:
        temporary.write_text(content, encoding="utf-8")
        temporary.replace(out_path)
    finally:
        if temporary.exists():
            temporary.unlink()
    return payload


def build_data(path=DEFAULT_XLSX, out_path=OUT_JS):
    payload = parse_sheet(path)
    write_output(payload, out_path)
    print(
        f"generated {payload['meta']['formationCount']} formations and "
        f"{payload['meta']['candidateCount']} candidates from {path}"
    )
    return payload


def main():
    parser = argparse.ArgumentParser(description="生成合阵浏览器数据")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--output", type=Path, default=OUT_JS)
    args = parser.parse_args()
    build_data(args.workbook, args.output)


if __name__ == "__main__":
    main()
