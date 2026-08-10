import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "秦时相关（更新贯侯钟离昧）20260618.xlsx"
OUT = ROOT / "data" / "inscription.js"
SHIELDS = ["天盾", "地盾", "人盾", "风盾", "云盾", "龙盾", "虎盾", "神盾", "鬼盾"]

def read_area(ws, start, end, quality):
    records = []
    current = None
    for row in range(start, end + 1):
        if ws.cell(row, 2).value:
            current = ws.cell(row, 2).value
        if not current:
            continue
        for index, shield in enumerate(SHIELDS):
            for col in (3 + index * 2, 4 + index * 2):
                name = ws.cell(row, col).value
                if name and str(name).strip() not in SHIELDS:
                    records.append({"name": str(name).strip(), "quality": quality, "tian": current, "shield": shield})
    return records

wb = load_workbook(BOOK, data_only=True, read_only=True)
ws = wb["铭文"]
records = read_area(ws, 3, 67, "普通橙色") + read_area(ws, 74, 133, "红色神将")
by_name = {}
for record in records:
    key = record["quality"] + "\0" + record["name"]
    by_name.setdefault(key, {"name": record["name"], "quality": record["quality"], "slots": []})["slots"].append({"tian": record["tian"], "shield": record["shield"]})
data = list(by_name.values())
OUT.write_text("window.INSCRIPTION_DATA = " + json.dumps({"items": data}, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print("generated", len(data), "disciples")
